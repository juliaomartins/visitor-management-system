"""The entrance report as a workbook.

SIX SHEETS, ONE QUESTION EACH, in the order somebody asks them: what happened,
when, where from, through which door, who is still missing, and finally the raw
log for anyone who wants to check the arithmetic.

The dump-everything-on-one-sheet approach is what makes an export boring — it
hands the reader the same work the system was supposed to do. Every sheet here
has already answered something.

EVERY SHEET OPENS WITH THE SHARED REPORT HEADER -- logo, navy title bar, event
name, dates -- from `apps/common/report_header.py`, the same header the visitor
and badge lists carry. A tab forwarded on its own still says what it is from.
So rows 1-3 belong to the header, headings sit on row 4 and data starts on row 5
on every table sheet; `HEADING_ROW` and `FIRST_ROW` below are the only place
that is written down.

Numbers are written as numbers, not strings, so a customer can pivot and chart
them without cleaning the file first. Percentages carry a number format rather
than a "%" glued onto text, for the same reason.
"""

import datetime as dt
import io

from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.common.report_header import HEADER_ROWS, write_xlsx_header
from apps.scans.models import ScanEvent, ScanResult

from .services import event_timezone

REPORT_TITLE = "ENTRANCE REPORT"

#: Rows 1-3 are the shared header. Table headings go on the next row.
HEADING_ROW = HEADER_ROWS + 1
FIRST_ROW = HEADING_ROW + 1

INK = "111827"
ACCENT = "2563EB"
HEADER_FILL = PatternFill("solid", fgColor=INK)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
LEAD_FONT = Font(size=10, color="4B5563")
KPI_FONT = Font(bold=True, size=22, color=ACCENT)
LABEL_FONT = Font(size=9, color="6B7280")
BODY = Font(size=10)
THIN = Side(style="thin", color="E5E7EB")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

REFUSED = {ScanResult.INVALID, ScanResult.REVOKED}

RESULT_LABEL = {
    ScanResult.VALID: "Valid",
    ScanResult.DUPLICATE: "Duplicate",
    ScanResult.REVOKED: "Revoked",
    ScanResult.INVALID: "Invalid",
}


def _headers(sheet, columns: list[tuple[str, int]], row: int = HEADING_ROW) -> None:
    for index, (label, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[row].height = 22
    # Freezing below the headings keeps the report header in view too.
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _table_sheet(workbook, name: str, columns: list[tuple[str, int]]):
    """A new tab: headings on HEADING_ROW, then the shared report header.

    The header goes in AFTER the headings because the headings set the column
    widths, and the header centres its logo in column A's final width.
    """
    sheet = workbook.create_sheet(name)
    _headers(sheet, columns)
    write_xlsx_header(sheet, REPORT_TITLE, last_column=len(columns))
    return sheet


def _row(sheet, row: int, values: list, percent_columns: tuple[int, ...] = ()) -> None:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.font = BODY
        cell.border = CELL_BORDER
        if index in percent_columns:
            cell.number_format = "0.0\\%"


def _local(stamp: dt.datetime, zone) -> dt.datetime:
    """openpyxl cannot write a tz-aware datetime, so drop the offset after
    converting. The sheet is labelled with the zone instead."""
    return timezone.localtime(stamp, zone).replace(tzinfo=None)


def build_workbook(
    queryset: QuerySet[ScanEvent],
    summary: dict,
    facts: dict,
    narrative: list[str],
    *,
    date_from: dt.date,
    date_to: dt.date,
) -> bytes:
    workbook = Workbook()
    zone = event_timezone()

    _summary_sheet(workbook, summary, facts, narrative, date_from, date_to, zone)
    _hourly_sheet(workbook, summary)
    _countries_sheet(workbook, summary)
    _doors_sheet(workbook, facts)
    _absent_sheet(workbook, facts)
    _log_sheet(workbook, queryset, zone)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary_sheet(workbook, summary, facts, narrative, date_from, date_to, zone):
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.column_dimensions["A"].width = 26
    for column in "BCDE":
        sheet.column_dimensions[column].width = 18

    write_xlsx_header(sheet, REPORT_TITLE, last_column=5)

    # What this report covers, under the header -- the header itself is the
    # same on every report, so the period has to be said here.
    span = (
        f"{date_from:%d %b %Y}"
        if date_from == date_to
        else f"{date_from:%d %b %Y} to {date_to:%d %b %Y}"
    )
    generated = timezone.localtime(timezone.now(), zone)
    meta = sheet.cell(
        row=HEADING_ROW,
        column=1,
        value=f"{span} · times in {zone} · Generated {generated:%d %b %Y %H:%M}",
    )
    meta.font = LEAD_FONT
    sheet.row_dimensions[HEADING_ROW].height = 20

    # The four figures that answer the question before anyone scrolls.
    kpi_label_row = HEADING_ROW + 2
    kpi_value_row = kpi_label_row + 1
    kpis = [
        ("Registered", facts["registered"]),
        ("Arrived", facts["arrived"]),
        ("Attendance", facts["attendance_rate"]),
        ("Refused", facts["refused"]),
    ]
    for index, (label, value) in enumerate(kpis):
        column = get_column_letter(1 + index)
        sheet[f"{column}{kpi_label_row}"] = label
        sheet[f"{column}{kpi_label_row}"].font = LABEL_FONT
        cell = sheet[f"{column}{kpi_value_row}"]
        cell.value = value
        cell.font = KPI_FONT
        if label == "Attendance":
            cell.number_format = "0.0\\%"
    sheet.row_dimensions[kpi_value_row].height = 30

    heading_row = kpi_value_row + 2
    sheet.cell(row=heading_row, column=1, value="What the numbers say").font = Font(
        bold=True, size=12, color=INK
    )

    for index, line in enumerate(narrative):
        row = heading_row + 1 + index
        cell = sheet.cell(row=row, column=1, value=f"•  {line}")
        cell.font = BODY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.row_dimensions[row].height = 30

    start = heading_row + 2 + len(narrative)
    sheet.cell(row=start, column=1, value="Outcome breakdown").font = Font(
        bold=True, size=12, color=INK
    )
    _headers(sheet, [("Outcome", 26), ("Scans", 14), ("Share", 14)], row=start + 1)
    sheet.freeze_panes = None  # a summary sheet does not scroll

    total = summary["total"] or 1
    for index, (result, label) in enumerate(RESULT_LABEL.items()):
        count = summary["by_result"][result]
        _row(
            sheet,
            start + 2 + index,
            [label, count, round(count / total * 100, 1)],
            percent_columns=(3,),
        )


def _hourly_sheet(workbook, summary):
    """Flow through the day, with a chart the customer does not have to build."""
    sheet = _table_sheet(
        workbook,
        "Hourly flow",
        [("Hour", 14), ("Total", 12), ("Valid", 12), ("Duplicate", 12), ("Refused", 12)],
    )

    for index, bucket in enumerate(summary["by_hour"]):
        hour = dt.datetime.fromisoformat(bucket["hour"])
        _row(
            sheet,
            FIRST_ROW + index,
            [
                f"{hour:%H:%M}",
                bucket["total"],
                bucket["valid"],
                bucket["duplicate"],
                bucket["refused"],
            ],
        )

    if not summary["by_hour"]:
        return

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Scans by hour"
    chart.y_axis.title = "Scans"
    chart.height = 8
    chart.width = 20

    last = HEADING_ROW + len(summary["by_hour"])
    # Columns C:E — valid, duplicate, refused — stacked. Not column B, which is
    # their sum and would double the height of every bar. The series names come
    # from HEADING_ROW, which is why these references start there and not at 1:
    # rows 1-3 are the report header, and a chart reading them would title its
    # series with the event name.
    chart.add_data(
        Reference(sheet, min_col=3, max_col=5, min_row=HEADING_ROW, max_row=last),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=FIRST_ROW, max_row=last))
    sheet.add_chart(chart, f"G{HEADING_ROW}")


def _countries_sheet(workbook, summary):
    sheet = _table_sheet(workbook, "Delegations", [("Country", 34), ("Scans", 14), ("Share", 14)])

    total = sum(row["total"] for row in summary["by_country"]) or 1
    for index, row in enumerate(summary["by_country"]):
        _row(
            sheet,
            FIRST_ROW + index,
            [row["country"] or "Unknown", row["total"], round(row["total"] / total * 100, 1)],
            percent_columns=(3,),
        )


def _doors_sheet(workbook, facts):
    sheet = _table_sheet(
        workbook,
        "Doors",
        [("Door", 30), ("Scans", 12), ("Admitted", 12), ("Refused", 12), ("Share", 12)],
    )

    for index, row in enumerate(facts["devices"]):
        _row(
            sheet,
            FIRST_ROW + index,
            [row["device"], row["total"], row["valid"], row["refused"], row["share"]],
            percent_columns=(5,),
        )


def _absent_sheet(workbook, facts):
    """Registered but never through the door — the call sheet."""
    sheet = _table_sheet(
        workbook,
        "Not arrived",
        [("Badge serial", 18), ("Full name", 30), ("Country", 22), ("Organisation", 30), ("Category", 12)],
    )

    for index, visitor in enumerate(facts["not_arrived"]):
        _row(
            sheet,
            FIRST_ROW + index,
            [
                visitor.badge_serial,
                visitor.full_name,
                visitor.country,
                visitor.organization or "",
                "VIP" if visitor.category == "vip" else "Normal",
            ],
        )


def _log_sheet(workbook, queryset, zone):
    """Every scan, so the arithmetic above can be checked."""
    sheet = _table_sheet(
        workbook,
        "Full log",
        [
            ("Scanned at", 20),
            ("Outcome", 14),
            ("Full name", 28),
            ("Country", 20),
            ("Organisation", 28),
            ("Category", 12),
            ("Badge serial", 18),
            ("Door", 22),
            ("Event ID", 12),
        ],
    )

    for index, scan in enumerate(queryset.iterator(chunk_size=500)):
        visitor = scan.visitor
        row = FIRST_ROW + index
        _row(
            sheet,
            row,
            [
                _local(scan.scanned_at, zone),
                RESULT_LABEL.get(scan.result, scan.result),
                visitor.full_name if visitor else "Unrecognised badge",
                visitor.country if visitor else "",
                (visitor.organization or "") if visitor else "",
                ("VIP" if visitor.category == "vip" else "Normal") if visitor else "",
                visitor.badge_serial if visitor else "",
                scan.device.name if scan.device else "",
                scan.id,
            ],
        )
        sheet.cell(row=row, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
