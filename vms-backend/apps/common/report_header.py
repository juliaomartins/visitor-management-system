"""The header every exported report carries, in Excel and in PDF.

ONE HEADER, ONE MODULE. The visitor list, the badge printing list and the
entrance report used to title themselves three different ways -- a navy banner
in two of them, a plain line in the third, a dark band across the PDF. They are
read side by side on the same desk, so they now open the same way:

    +------+----------------------------------------------+
    |      | ########   TITLE OF THIS REPORT   ########### |   navy bar
    | LOGO |      DRCC AND MINISTERIAL DIALOGUE 2026      |   bold italic
    |      |     2-3 OCTOBER 2026 . DILI, TIMOR-LESTE     |   grey italic
    +------+----------------------------------------------+

The event name, the dates, the colour and the logo live here and nowhere else,
so a report cannot drift from its siblings. Only the title belongs to the caller.

NOT TRANSLATED, for the reason the printed badge is not: these files leave the
dashboard and are read by whoever receives them, and the event's official title
is not ours to translate.

THE LOGO SITS ON WHITE, NEVER ON THE NAVY BAR. It is a full-colour emblem with a
blue plume and a black ribbon, and both disappear against #00309C. So it gets a
cell (Excel) or a margin (PDF) of its own, left of all three lines.

`brand/logo.png` is a copy of `vms-dashboard/public/brand/logo.png`, trimmed of
its transparent border and scaled to 320 px -- the same duplication, for the same
reason, as `apps/badges/assets/drcc-event.png`: a backend-only deploy does not
check out the dashboard's `public/` folder. Replace both copies together.
"""

import io
from functools import lru_cache
from pathlib import Path

from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from PIL import Image as PilImage
from reportlab.lib import colors
from reportlab.lib.units import mm

EVENT_NAME = "DRCC AND MINISTERIAL DIALOGUE 2026"
EVENT_WHEN = "2–3 OCTOBER 2026 · DÍLI, TIMOR-LESTE"

# The event's own blue, sampled from the RDTL and SECoop logos -- the same
# `--color-brand-blue` the dashboard, the badge and the lobby screen use.
BRAND_BLUE = "00309C"
INK = "111827"
MUTED = "4A5560"
RULE = "D9DDE3"

LOGO_PATH = Path(__file__).resolve().parent / "brand" / "logo.png"


# --------------------------------------------------------------------- Excel --

#: The header occupies rows 1-3; a sheet's own content starts at row 4.
HEADER_ROWS = 3

_TITLE_FILL = PatternFill("solid", fgColor=BRAND_BLUE)
_LINES = (
    # (font, row height in points)
    (Font(color="FFFFFF", bold=True, size=14), 30),
    (Font(bold=True, italic=True, size=11, color=INK), 20),
    (Font(italic=True, size=10, color=MUTED), 18),
)

#: Column A is widened to at least this, in Excel character units (~103 px),
#: so the logo has a white cell of its own beside the three lines.
LOGO_COLUMN_WIDTH = 14
#: Displayed height of the logo in the sheet. The embedded PNG is twice this,
#: so it stays sharp when the sheet is zoomed or printed.
LOGO_XLSX_PX = 80

#: A merged cell does not let its text overflow -- it clips. The title and
#: date lines need about this many characters of width, so a sheet whose table
#: is narrower than that gets a banner that runs on past its last column.
MIN_BANNER_CHARS = 60
#: What Excel draws a column that was never given a width.
_DEFAULT_COLUMN_CHARS = 8.43


@lru_cache(maxsize=1)
def _logo_for_xlsx() -> tuple[bytes, tuple[int, int]]:
    """The logo at 2x its display height, as PNG bytes, and its display size.

    PNG and not JPEG: the emblem is flat colour on transparency, which JPEG
    would ring around and fill with a white box.
    """
    with PilImage.open(LOGO_PATH) as source:
        image = source.convert("RGBA")
    height = LOGO_XLSX_PX
    width = round(height * image.width / image.height)
    image = image.resize((width * 2, height * 2), PilImage.LANCZOS)

    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue(), (width, height)


def _column_chars(sheet, column: int) -> float:
    """A column's width in characters, without creating a dimension for it.

    `column_dimensions` is a defaulting dictionary: reading a missing key adds
    one, with openpyxl's default of 13 -- wider than what Excel actually draws.
    """
    letter = get_column_letter(column)
    if letter in sheet.column_dimensions and sheet.column_dimensions[letter].width:
        return sheet.column_dimensions[letter].width
    return _DEFAULT_COLUMN_CHARS


def _banner_last_column(sheet, last_column: int) -> int:
    """The last column the text lines span: the table's, or further if narrow."""
    column, total = 2, 0.0
    while True:
        total += _column_chars(sheet, column)
        if column >= last_column and total >= MIN_BANNER_CHARS:
            return column
        column += 1


def write_xlsx_header(sheet, title: str, *, last_column: int) -> None:
    """Write the report header into rows 1-3 of `sheet`.

    `last_column` is the table's last column; the banner spans at least that.

    CALL IT AFTER THE SHEET'S COLUMN WIDTHS ARE SET. The logo is centred in
    column A's final width and the banner is sized from the widths it can see;
    widths written afterwards leave a logo off-centre or a date line clipped.

    STYLE ONLY THE ANCHOR OF A MERGED RANGE. `Worksheet._clean_merge_range`
    replaces every cell but the top-left with a fresh `MergedCell` and restores
    borders only, because Excel draws a merged range from the top-left cell's
    fill, font and alignment. A fill written to the others is discarded.
    """
    if _column_chars(sheet, 1) < LOGO_COLUMN_WIDTH:
        sheet.column_dimensions["A"].width = LOGO_COLUMN_WIDTH

    last = _banner_last_column(sheet, max(last_column, 2))

    sheet.merge_cells(start_row=1, start_column=1, end_row=HEADER_ROWS, end_column=1)

    for row, (text, (font, height)) in enumerate(
        zip((title, EVENT_NAME, EVENT_WHEN), _LINES), start=1
    ):
        if last > 2:
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last)
        cell = sheet.cell(row=row, column=2, value=text)
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if row == 1:
            cell.fill = _TITLE_FILL
        sheet.row_dimensions[row].height = height

    _place_logo(sheet)


def _place_logo(sheet) -> None:
    """Pin the logo, centred, inside the merged A1:A3 cell.

    A one-cell anchor with explicit offsets rather than the string anchor "A1":
    the string form puts the image flush in the top-left corner, touching the
    grid lines, which reads as a mistake rather than a header.
    """
    png, (width, height) = _logo_for_xlsx()

    column_px = int(_column_chars(sheet, 1) * 7 + 5)
    rows_px = sum(points_to_pixels(line_height) for _, line_height in _LINES)

    image = XlsxImage(io.BytesIO(png))
    image.width, image.height = width, height
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=0,
            colOff=pixels_to_EMU(max(0, (column_px - width) // 2)),
            row=0,
            rowOff=pixels_to_EMU(max(0, (rows_px - height) // 2)),
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
    )
    sheet.add_image(image)


# ----------------------------------------------------------------------- PDF --

_TOP_GAP = 10 * mm
_LOGO_H = 22 * mm
_BAR_H = 8.5 * mm
#: From the top edge of the page to the bottom of the meta line. A document's
#: body frame starts below this.
PDF_HEADER_HEIGHT = 41 * mm


def draw_pdf_header(
    canvas,
    *,
    page_width: float,
    page_height: float,
    margin: float,
    title: str,
    meta_left: str = "",
    meta_right: str = "",
) -> None:
    """Draw the report header at the top of the current page.

    Meant to be called from a page template's `onPage`, so every page carries
    it. White with one navy bar, so a report printed on an office printer does
    not spend a page's worth of toner on a band.

    `meta_left` / `meta_right` sit in small grey type under a hairline -- the
    period covered and when the file was made -- so the three event lines above
    stay exactly the same on every report.
    """
    canvas.saveState()

    top = page_height - _TOP_GAP

    # The logo, in its own white space at the left of all three lines.
    logo_w = _LOGO_H * 304 / 320
    canvas.drawImage(
        str(LOGO_PATH),
        margin,
        top - _LOGO_H,
        width=logo_w,
        height=_LOGO_H,
        mask="auto",
        preserveAspectRatio=True,
    )

    left = margin + logo_w + 5 * mm
    right = page_width - margin
    centre = (left + right) / 2

    canvas.setFillColor(colors.HexColor(f"#{BRAND_BLUE}"))
    canvas.rect(left, top - _BAR_H, right - left, _BAR_H, stroke=0, fill=1)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 12.5)
    canvas.drawCentredString(centre, top - 5.9 * mm, title)

    canvas.setFillColor(colors.HexColor(f"#{INK}"))
    canvas.setFont("Helvetica-BoldOblique", 10)
    canvas.drawCentredString(centre, top - 14 * mm, EVENT_NAME)

    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.setFont("Helvetica-Oblique", 8.8)
    canvas.drawCentredString(centre, top - 19.5 * mm, EVENT_WHEN)

    rule_y = top - 25 * mm
    canvas.setStrokeColor(colors.HexColor(f"#{RULE}"))
    canvas.setLineWidth(0.6)
    canvas.line(margin, rule_y, right, rule_y)

    canvas.setFont("Helvetica", 7.8)
    if meta_left:
        canvas.drawString(margin, rule_y - 4.5 * mm, meta_left)
    if meta_right:
        canvas.drawRightString(right, rule_y - 4.5 * mm, meta_right)

    canvas.restoreState()
